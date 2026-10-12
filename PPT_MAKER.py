"""
ppt_note_toolkit_gui.py

PPT 발표자 노트 & 변환 도구 모음 (GUI, 탭 4개)
------------------------------------------------------------
① 노트 삽입   : 마크다운(.md) 또는 텍스트(.txt) 파일의 노트를 기존 PPTX의
                각 슬라이드 '노트' 란에 순서대로 삽입합니다.
② 노트 추출   : 기존 PPTX에서 발표자 노트를 꺼내 텍스트 파일로 저장합니다.
                (마크다운 형식 / 구분선(-----) 형식 중 선택)
③ 음성 변환   : ①·②와 동일한 노트 파일을 읽어 Microsoft Azure AI Speech로
                슬라이드별 자연스러운 MP3 음성 파일을 생성합니다. (하나로 합치기도 가능)
                Azure AI Speech는 매달 뉴럴 TTS 50만 자까지 무료(F0 티어)입니다.
④ PDF → PPT  : PDF의 각 페이지를 고해상도 JPG로 렌더링한 뒤, 슬라이드 전체에
                꽉 채운 PPTX로 저장합니다. (이미지 삽입 방식이라 폰트/레이아웃이
                절대 깨지지 않습니다)

설치
------------------------------------------------------------
    pip install python-pptx                    # 필수 (①②③④ 공통)
    pip install azure-cognitiveservices-speech  # 선택 (③ 음성 변환을 쓸 때만 필요)
    pip install pymupdf pillow                  # 선택 (④ PDF → PPT 를 쓸 때만 필요)

실행
------------------------------------------------------------
    python ppt_note_toolkit_gui.py

지원하는 노트 파일 형식 (자동 인식)
------------------------------------------------------------
1) 마크다운 형식 — PPT 재삽입 시 가장 안전(권장)
       ## 슬라이드 01 (표지)
       **슬라이드 노트 내용**
       안녕하십니까 ...

       ---

       ## 슬라이드 02
       **슬라이드 노트 내용**
       당 프로젝트의 ...

2) 구분선 형식 — 내용이 단순할 때, 사람이 읽고 편집하기 편함
       안녕하십니까 ...

       -----

       당 프로젝트의 ...

       -----
   "-----" (하이픈 3개 이상)만 있는 줄을 슬라이드 경계로 인식하며,
   슬라이드 번호는 등장 순서(1, 2, 3 ...)로 자동 부여됩니다.

3) 레거시 텍스트 형식(구버전 호환) — 한 줄 = 슬라이드 하나
   위 두 형식에 해당하지 않을 때만 사용됩니다.

음성 변환(③) 관련 참고
------------------------------------------------------------
- Azure Portal(portal.azure.com)에서 "Speech" 리소스를 만들면 API 키와 리전이
  발급됩니다. F0(무료) 가격 책정 티어를 선택하면 매달 뉴럴 TTS 50만 자까지
  과금 없이 사용할 수 있습니다.
- "음성 목록 불러오기" 버튼을 누르면 입력한 언어(로케일) 기준으로 사용 가능한
  뉴럴 음성들이 드롭다운에 채워집니다. voice 이름(예: ko-KR-SunHiNeural)을
  이미 알고 있다면 직접 입력해도 됩니다.
- 슬라이드별로 mp3를 개별 생성한 뒤, 선택 시 하나의 파일로 이어붙입니다.
  단, 이어붙이기는 재인코딩 없는 단순 결합이라 완벽하게 매끄럽지 않을 수
  있습니다. (ffmpeg 등 추가 프로그램 설치 없이 동작하도록 설계했습니다)
"""

import os
import platform
import re
import subprocess
import sys
import tempfile
import threading
import traceback
from pathlib import Path

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

try:
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.util import Emu
except ImportError:
    print(
        "python-pptx 가 설치되어 있지 않습니다.\n"
        "터미널에서 아래 명령을 실행한 뒤 다시 실행하세요.\n\n"
        "    pip install python-pptx\n"
    )
    sys.exit(1)

# pymupdf / pillow는 ④ PDF → PPT 탭에서만 필요한 선택 의존성이라, 미설치 상태여도
# 앱 자체는 정상 실행되도록 하고 해당 탭을 사용할 때만 안내 메시지를 띄웁니다.
try:
    import fitz  # PyMuPDF
    from PIL import Image
    HAVE_PDF_LIBS = True
except ImportError:
    HAVE_PDF_LIBS = False


# ══════════════════════════════════════════════════════════════
# 공통: 노트 파일 파싱 / 포맷 변환
# ══════════════════════════════════════════════════════════════

_MD_HEADER_RE = re.compile(r"(?m)^##\s*슬라이드\s*(\d+)")
_DASH_LINE_RE = re.compile(r"(?m)^\s*-{3,}\s*$")


def _parse_markdown_notes(text):
    blocks = re.split(r"(?m)^##\s*슬라이드\s*(\d+)", text)
    notes, order = {}, []
    for i in range(1, len(blocks), 2):
        slide_no = int(blocks[i])
        block_text = blocks[i + 1]
        m = re.search(r"\*\*슬라이드 노트 내용\*\*\s*\n(.+?)(?:\n-{3,}|\Z)", block_text, re.S)
        note_text = (m.group(1).strip() if m else block_text.strip())
        notes[slide_no] = note_text
        order.append(slide_no)
    return notes, order


def _parse_dash_separated_notes(text):
    parts = _DASH_LINE_RE.split(text)
    notes, order = {}, []
    for idx, part in enumerate(parts, start=1):
        content = part.strip()
        if content == "" and idx == len(parts):
            continue  # 파일이 구분선으로 끝나서 생기는 빈 마지막 조각은 무시
        notes[idx] = content
        order.append(idx)
    return notes, order


def _parse_plain_text_notes(text):
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    notes, order = {}, []
    for idx, line in enumerate(lines, start=1):
        notes[idx] = line
        order.append(idx)
    return notes, order


def load_notes(notes_file_path):
    """마크다운 / 구분선 / 레거시(줄당 1슬라이드) 형식을 자동 판별해 파싱합니다."""
    with open(notes_file_path, "r", encoding="utf-8-sig") as f:
        text = f.read()

    if _MD_HEADER_RE.search(text):
        notes, order = _parse_markdown_notes(text)
        fmt = "markdown"
    elif _DASH_LINE_RE.search(text):
        notes, order = _parse_dash_separated_notes(text)
        fmt = "dash-separated"
    else:
        notes, order = _parse_plain_text_notes(text)
        fmt = "plain-text(legacy)"

    return notes, order, fmt


def format_notes_as_markdown(notes_data):
    """notes_data: [(slide_no, note_text), ...] → 마크다운 문자열"""
    parts = []
    for slide_no, note in notes_data:
        note_text = note.strip() if note and note.strip() else "(발표자 노트 없음)"
        parts.append(f"## 슬라이드 {slide_no:02d}\n**슬라이드 노트 내용**\n{note_text}\n\n---\n")
    return "\n".join(parts)


def format_notes_as_dash(notes_data, dash="-----"):
    """notes_data: [(slide_no, note_text), ...] → 구분선 문자열 (슬라이드 번호 표기 없음)"""
    lines = []
    for idx, (slide_no, note) in enumerate(notes_data):
        if idx > 0:
            lines.append("")
            lines.append(dash)
            lines.append("")
        lines.append(note.strip() if note and note.strip() else "(발표자 노트 없음)")
    return "\n".join(lines) + "\n"


# ══════════════════════════════════════════════════════════════
# ① 노트 삽입 (텍스트 → PPT)
# ══════════════════════════════════════════════════════════════

_NOTES_PH_LEFT = Emu(685800)
_NOTES_PH_TOP = Emu(4343400)
_NOTES_PH_WIDTH = Emu(5486400)
_NOTES_PH_HEIGHT = Emu(4114800)


def ensure_notes_placeholder(notes_slide):
    """노트 슬라이드에 '노트 본문' 플레이스홀더가 없으면 새로 만들어줍니다."""
    placeholder = notes_slide.notes_placeholder
    if placeholder is not None:
        return placeholder

    spTree = notes_slide.shapes._spTree
    next_id = spTree.max_shape_id + 1
    spTree.add_placeholder(next_id, "Notes Placeholder", PP_PLACEHOLDER.BODY, "horz", "quarter", 1)

    placeholder = notes_slide.notes_placeholder
    placeholder.left = _NOTES_PH_LEFT
    placeholder.top = _NOTES_PH_TOP
    placeholder.width = _NOTES_PH_WIDTH
    placeholder.height = _NOTES_PH_HEIGHT
    return placeholder


def check_pptx_path(path, log):
    log(f"입력 경로: {path}")
    log(f"절대 경로: {os.path.abspath(path)}")

    if path != path.strip():
        log("⚠ 경로 앞/뒤에 공백이 포함되어 있습니다.")

    if not os.path.exists(path):
        log("✗ 이 경로에 파일(또는 폴더)이 존재하지 않습니다.")
        folder = os.path.dirname(path)
        target = os.path.basename(path)
        if os.path.isdir(folder):
            log(f"  → 폴더 '{folder}' 안의 실제 파일 목록:")
            for name in os.listdir(folder):
                marker = "  ← 이름이 비슷함" if name.lower().replace(" ", "") == target.lower().replace(" ", "") else ""
                log(f"    - {name}{marker}")
        else:
            log(f"  → 상위 폴더 '{folder}' 자체가 존재하지 않습니다.")
        return False

    if os.path.isdir(path):
        log("✗ 이 경로는 파일이 아니라 폴더입니다.")
        return False

    size = os.path.getsize(path)
    log(f"파일 크기: {size:,} bytes")
    if size == 0:
        log(
            "✗ 파일 크기가 0바이트입니다. OneDrive '온라인 전용' 상태일 수 있습니다.\n"
            "  → 탐색기에서 파일 우클릭 → '이 디바이스에 항상 유지' 선택 후 다시 시도하세요."
        )
        return False

    if not path.lower().endswith(".pptx"):
        log("⚠ 확장자가 .pptx가 아닙니다.")

    log("✓ 경로/파일 상태에는 문제가 없어 보입니다.")
    return True


def insert_slide_notes_from_file(notes_file_path, input_pptx_path, output_pptx_path, dry_run, log):
    if not check_pptx_path(input_pptx_path, log):
        log("\n위 문제를 해결한 뒤 다시 실행하세요. (아직 노트는 삽입되지 않았습니다)")
        return False

    try:
        prs = Presentation(input_pptx_path)
    except FileNotFoundError:
        log(f"오류: 원본 PPTX '{input_pptx_path}' 파일을 찾을 수 없습니다.")
        return False
    except Exception as e:
        log(f"오류: PPTX 파일을 여는 중 문제가 발생했습니다: {type(e).__name__}: {e}")
        log("  → 파일이 손상되었거나 PowerPoint에서 열려 있어 잠겨 있을 수 있습니다.")
        return False

    try:
        notes, order, fmt = load_notes(notes_file_path)
    except FileNotFoundError:
        log(f"오류: 노트 파일 '{notes_file_path}' 을(를) 찾을 수 없습니다.")
        return False

    if not notes:
        log("오류: 노트 파일에서 슬라이드 노트를 하나도 찾지 못했습니다.")
        return False

    n_slides = len(prs.slides)
    n_notes = len(order)
    log(f"노트 파일 형식: {fmt}")
    log(f"PPT 슬라이드 수: {n_slides}장 / 노트 개수: {n_notes}개")

    if n_slides != n_notes:
        log("⚠ 슬라이드 수와 노트 개수가 다릅니다. 순서(위치) 기준으로 매칭합니다. 결과를 반드시 확인하세요.")

    applied, skipped = 0, 0
    for position, slide_no in enumerate(order):
        note_text = notes[slide_no]

        if position >= n_slides:
            log(f"  건너뜀 → PPT에 {position + 1}번째 슬라이드가 없음 (노트 원본: 슬라이드 {slide_no})")
            skipped += 1
            continue

        if dry_run:
            preview = note_text.replace("\n", " ")[:50]
            log(f"  [미리보기] PPT {position + 1}번 슬라이드 ← 노트파일 슬라이드 {slide_no} : {preview}...")
            continue

        slide = prs.slides[position]
        try:
            notes_slide = slide.notes_slide
            ensure_notes_placeholder(notes_slide)
            notes_slide.notes_text_frame.text = note_text
            applied += 1
        except Exception as e:
            log(f"  ✗ PPT {position + 1}번 슬라이드 노트 삽입 실패: {type(e).__name__}: {e}")
            skipped += 1

    if dry_run:
        log("\n[미리보기 모드] 실제 저장은 하지 않았습니다.")
        return True

    try:
        prs.save(output_pptx_path)
    except Exception as e:
        log(f"오류: 저장 중 문제가 발생했습니다: {type(e).__name__}: {e}")
        return False

    log(f"\n총 {applied}개 슬라이드에 노트를 삽입했습니다. (건너뜀 {skipped}개)")
    log(f"'{output_pptx_path}' 파일로 저장했습니다. (원본 파일은 수정하지 않았습니다)")
    return True


# ══════════════════════════════════════════════════════════════
# ② 노트 추출 (PPT → 텍스트)
# ══════════════════════════════════════════════════════════════

def extract_notes_from_pptx(input_pptx_path, log):
    """PPTX의 모든 슬라이드 노트를 [(슬라이드번호, 노트텍스트), ...] 로 반환합니다."""
    prs = Presentation(input_pptx_path)
    notes_data = []
    for i, slide in enumerate(prs.slides, start=1):
        text = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            text = slide.notes_slide.notes_text_frame.text
        notes_data.append((i, text.strip()))
    empty_count = sum(1 for _, t in notes_data if not t)
    log(f"총 {len(notes_data)}장의 슬라이드에서 노트를 추출했습니다. (노트 없는 슬라이드 {empty_count}장)")
    return notes_data


# ══════════════════════════════════════════════════════════════
# ③ 음성 변환 (텍스트 → MP3, Microsoft Azure AI Speech)
# ══════════════════════════════════════════════════════════════

_TTS_MAX_CHARS = 4500  # 1회 요청에 넉넉한 여유를 둔 청크 크기
_AZURE_OUTPUT_FORMAT_NAME = "Audio24Khz160KBitRateMonoMp3"  # 고음질 mp3 출력 포맷

# 리전 입력칸에 참고용으로 보여줄 자주 쓰는 리전 후보 (직접 다른 값을 입력해도 됩니다)
AZURE_COMMON_REGIONS = ["koreacentral", "japaneast", "eastus", "eastus2", "westus", "southeastasia"]
AZURE_DEFAULT_LOCALE = "ko-KR"


def _split_text_for_tts(text, max_chars=_TTS_MAX_CHARS):
    text = text.strip()
    if len(text) <= max_chars:
        return [text]
    sentences = re.split(r"(?<=[.!?\n])\s+", text)
    chunks, current = [], ""
    for s in sentences:
        candidate = f"{current} {s}".strip() if current else s
        if len(candidate) > max_chars and current:
            chunks.append(current.strip())
            current = s
        else:
            current = candidate
    if current:
        chunks.append(current.strip())
    return chunks


def concatenate_mp3_files(file_paths, output_path):
    """여러 mp3 파일을 순서대로 이어붙입니다. (재인코딩 없는 단순 결합)"""
    with open(output_path, "wb") as out_f:
        for p in file_paths:
            with open(p, "rb") as in_f:
                out_f.write(in_f.read())


def _extract_voice_name(selection):
    """
    '이름 (voice_short_name)' 형태로 표시되는 콤보박스 선택값에서 voice 이름만 뽑아냅니다.
    괄호가 없으면(= 사용자가 voice 이름을 직접 입력한 경우) 입력값 그대로를 사용합니다.
    """
    selection = (selection or "").strip()
    m = re.search(r"\(([^()]+)\)\s*$", selection)
    return m.group(1).strip() if m else selection


def list_azure_voices(api_key, region, locale_filter, log):
    """지정한 리전/로케일에서 사용 가능한 음성 목록을 [(표시이름, short_name), ...] 로 반환합니다."""
    try:
        import azure.cognitiveservices.speech as speechsdk
    except ImportError:
        log("오류: azure-cognitiveservices-speech 패키지가 설치되어 있지 않습니다. 'pip install azure-cognitiveservices-speech' 를 실행하세요.")
        return []

    if not api_key or not region:
        log("오류: API 키와 리전을 모두 입력하세요.")
        return []

    try:
        speech_config = speechsdk.SpeechConfig(subscription=api_key, region=region)
        synthesizer = speechsdk.SpeechSynthesizer(speech_config=speech_config, audio_config=None)
        result = synthesizer.get_voices_async(locale_filter or "").get()

        if result.reason != speechsdk.ResultReason.VoicesListRetrieved:
            log(f"음성 목록을 불러오지 못했습니다: {result.reason}")
            return []

        voices = []
        for v in result.voices:
            gender = getattr(v.gender, "name", str(v.gender))
            voices.append((f"{v.local_name} · {gender} ({v.locale})", v.short_name))
        log(f"사용 가능한 음성 {len(voices)}개를 불러왔습니다.")
        return voices
    except Exception as e:
        log(f"음성 목록을 불러오지 못했습니다: {type(e).__name__}: {e}")
        return []


def _tts_request_to_file(api_key, region, voice_name, text, output_path):
    import azure.cognitiveservices.speech as speechsdk

    speech_config = speechsdk.SpeechConfig(subscription=api_key, region=region)
    speech_config.speech_synthesis_voice_name = voice_name
    speech_config.set_speech_synthesis_output_format(
        getattr(speechsdk.SpeechSynthesisOutputFormat, _AZURE_OUTPUT_FORMAT_NAME)
    )

    audio_config = speechsdk.audio.AudioOutputConfig(filename=output_path)
    synthesizer = speechsdk.SpeechSynthesizer(speech_config=speech_config, audio_config=audio_config)
    result = synthesizer.speak_text_async(text).get()

    if result.reason == speechsdk.ResultReason.SynthesizingAudioCompleted:
        return
    if result.reason == speechsdk.ResultReason.Canceled:
        details = result.cancellation_details
        raise RuntimeError(f"{details.reason}: {details.error_details}")
    raise RuntimeError(f"음성 합성에 실패했습니다 (reason={result.reason})")


def synthesize_slide_audio(api_key, region, voice_name, text, output_path):
    chunks = _split_text_for_tts(text)
    if len(chunks) == 1:
        _tts_request_to_file(api_key, region, voice_name, chunks[0], output_path)
        return

    tmp_dir = tempfile.mkdtemp(prefix="tts_chunk_")
    tmp_paths = []
    try:
        for i, chunk in enumerate(chunks, start=1):
            tmp_path = os.path.join(tmp_dir, f"part_{i:03d}.mp3")
            _tts_request_to_file(api_key, region, voice_name, chunk, tmp_path)
            tmp_paths.append(tmp_path)
        concatenate_mp3_files(tmp_paths, output_path)
    finally:
        for p in tmp_paths:
            try:
                os.remove(p)
            except OSError:
                pass
        try:
            os.rmdir(tmp_dir)
        except OSError:
            pass


def parse_slide_range(range_str, available_slide_numbers):
    """'1-5,10,12-14' 같은 범위 문자열을 슬라이드 번호 set으로 변환합니다. 빈 문자열이면 None(=전체)."""
    range_str = (range_str or "").strip()
    if not range_str:
        return None

    selected = set()
    for part in range_str.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            try:
                a, b = int(a.strip()), int(b.strip())
            except ValueError:
                continue
            selected.update(range(min(a, b), max(a, b) + 1))
        else:
            try:
                selected.add(int(part))
            except ValueError:
                continue
    return selected & set(available_slide_numbers)


def generate_audio_for_notes(notes_data, api_key, region, voice_selection, output_dir, combine, combined_path, slide_filter, log):
    try:
        import azure.cognitiveservices.speech  # noqa: F401  (설치 여부만 확인)
    except ImportError:
        log("오류: azure-cognitiveservices-speech 패키지가 설치되어 있지 않습니다. 터미널에서 'pip install azure-cognitiveservices-speech' 를 실행하세요.")
        return False

    if not api_key or not region:
        log("오류: API 키와 리전을 모두 입력하세요.")
        return False

    voice_name = _extract_voice_name(voice_selection)
    if not voice_name:
        log("오류: 사용할 음성을 선택하거나 voice 이름을 입력하세요.")
        return False

    os.makedirs(output_dir, exist_ok=True)

    target_slides = notes_data if slide_filter is None else [(n, t) for n, t in notes_data if n in slide_filter]
    if not target_slides:
        log("변환할 슬라이드가 없습니다. 슬라이드 범위를 확인하세요.")
        return False

    log(f"총 {len(target_slides)}개 슬라이드 음성 생성을 시작합니다. (리전: {region}, 음성: {voice_name})")

    generated_paths = []
    for slide_no, text in target_slides:
        if not text.strip():
            log(f"  슬라이드 {slide_no}: 노트 내용이 없어 건너뜁니다.")
            continue
        out_path = os.path.join(output_dir, f"slide_{slide_no:03d}.mp3")
        log(f"  슬라이드 {slide_no} 음성 생성 중... ({len(text)}자)")
        try:
            synthesize_slide_audio(api_key, region, voice_name, text, out_path)
            generated_paths.append(out_path)
            log(f"    → 저장: {out_path}")
        except Exception as e:
            log(f"    ✗ 슬라이드 {slide_no} 음성 생성 실패: {type(e).__name__}: {e}")

    log(f"\n총 {len(generated_paths)}개 슬라이드의 음성 파일을 생성했습니다.")

    if combine and generated_paths:
        try:
            concatenate_mp3_files(generated_paths, combined_path)
            log(f"모든 음성을 하나로 합쳐 '{combined_path}' 로 저장했습니다.")
            log("※ 단순 결합 방식이라 슬라이드 사이 무음 간격 등 매끄러움은 보장되지 않습니다. 필요하면 오디오 편집기로 다듬어주세요.")
        except Exception as e:
            log(f"오류: 음성 파일 합치기 실패: {type(e).__name__}: {e}")

    return True


# ══════════════════════════════════════════════════════════════
# ④ PDF → PPT (PDF 페이지 → 고화질 JPG → 이미지 꽉 채운 PPTX)
# ══════════════════════════════════════════════════════════════

_PDF_EMU_PER_PX = 9525  # 96dpi 기준 1px = 9525 EMU


def pdf_to_images(pdf_path, out_dir, dpi, quality, progress_cb=None):
    """PDF의 각 페이지를 고해상도 JPG로 렌더링해 out_dir에 저장하고 경로 목록을 반환합니다."""
    doc = fitz.open(pdf_path)
    zoom = dpi / 72
    matrix = fitz.Matrix(zoom, zoom)
    total = len(doc)
    image_paths = []
    try:
        for i, page in enumerate(doc):
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            img_path = os.path.join(out_dir, f"page_{i + 1:03d}.jpg")
            pix.save(img_path, jpg_quality=quality)
            image_paths.append(img_path)
            if progress_cb:
                progress_cb(i + 1, total, f"[JPG] {i + 1}/{total} 페이지 저장 완료 ({pix.width}x{pix.height}px)")
    finally:
        doc.close()
    return image_paths


def images_to_ppt(image_paths, output_pptx, progress_cb=None):
    """렌더링된 이미지들을 슬라이드 전체에 꽉 채운 PPTX로 저장합니다."""
    if not image_paths:
        raise ValueError("변환할 이미지가 없습니다.")

    with Image.open(image_paths[0]) as im:
        px_w, px_h = im.size

    prs = Presentation()
    prs.slide_width = Emu(px_w * _PDF_EMU_PER_PX)
    prs.slide_height = Emu(px_h * _PDF_EMU_PER_PX)
    blank_layout = prs.slide_layouts[6]

    total = len(image_paths)
    for i, img_path in enumerate(image_paths):
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(img_path, left=0, top=0, width=prs.slide_width, height=prs.slide_height)
        if progress_cb:
            progress_cb(i + 1, total, f"[PPT] {i + 1}/{total} 슬라이드 생성 완료")

    prs.save(output_pptx)


def convert_pdf_to_pptx(pdf_path, output_pptx, dpi, quality, progress_cb, log):
    """PDF → JPG → PPTX 전체 과정을 순서대로 실행합니다."""
    if not HAVE_PDF_LIBS:
        log(
            "오류: pymupdf / pillow 가 설치되어 있지 않습니다.\n"
            "  터미널에서 아래 명령을 실행한 뒤 다시 시도하세요.\n\n"
            "    pip install pymupdf pillow\n"
        )
        return False

    if not os.path.exists(pdf_path):
        log(f"오류: PDF 파일 '{pdf_path}' 을(를) 찾을 수 없습니다.")
        return False

    base_name = os.path.splitext(os.path.basename(pdf_path))[0]
    out_dir = os.path.join(os.path.dirname(output_pptx) or ".", f"{base_name}_slides")
    os.makedirs(out_dir, exist_ok=True)

    log(f"PDF → JPG 변환 중... (DPI {dpi}, 품질 {quality}, 임시 폴더: {out_dir})")
    try:
        image_paths = pdf_to_images(pdf_path, out_dir, dpi, quality, progress_cb=progress_cb)
    except Exception as e:
        log(f"오류: PDF를 이미지로 변환하는 중 문제가 발생했습니다: {type(e).__name__}: {e}")
        return False

    log("PPT 생성 중...")
    try:
        images_to_ppt(image_paths, output_pptx, progress_cb=progress_cb)
    except Exception as e:
        log(f"오류: PPTX 저장 중 문제가 발생했습니다: {type(e).__name__}: {e}")
        return False

    log(f"\n✅ 변환 완료: {output_pptx}")
    return True


def open_in_file_manager(path):
    folder = os.path.dirname(path) or "."
    system = platform.system()
    try:
        if system == "Windows":
            os.startfile(folder)
        elif system == "Darwin":
            subprocess.run(["open", folder])
        else:
            subprocess.run(["xdg-open", folder])
    except Exception:
        pass


# ══════════════════════════════════════════════════════════════
# GUI
# ══════════════════════════════════════════════════════════════

class InsertTab(ttk.Frame):
    def __init__(self, parent, log_fn, shared_state):
        super().__init__(parent)
        self.log = log_fn
        self.shared_state = shared_state

        self.notes_path = tk.StringVar()
        self.input_pptx_path = tk.StringVar()
        self.output_pptx_path = tk.StringVar()

        self._build()

    def _build(self):
        pad = {"padx": 10, "pady": 6}
        form = ttk.Frame(self)
        form.pack(fill="x", **pad)
        form.columnconfigure(0, weight=1)

        ttk.Label(form, text="① 슬라이드 노트 파일 (.md / .txt)").grid(row=0, column=0, columnspan=2, sticky="w")
        ttk.Entry(form, textvariable=self.notes_path).grid(row=1, column=0, sticky="ew", padx=(0, 8))
        ttk.Button(form, text="찾아보기", command=self._browse_notes).grid(row=1, column=1)

        ttk.Label(form, text="② 원본 PPTX 파일").grid(row=2, column=0, columnspan=2, sticky="w", pady=(14, 0))
        ttk.Entry(form, textvariable=self.input_pptx_path).grid(row=3, column=0, sticky="ew", padx=(0, 8))
        ttk.Button(form, text="찾아보기", command=self._browse_pptx).grid(row=3, column=1)

        ttk.Label(form, text="③ 저장할 파일 (노트가 삽입된 새 PPTX)").grid(row=4, column=0, columnspan=2, sticky="w", pady=(14, 0))
        ttk.Entry(form, textvariable=self.output_pptx_path).grid(row=5, column=0, sticky="ew", padx=(0, 8))
        ttk.Button(form, text="다른 이름으로 저장", command=self._browse_output).grid(row=5, column=1)

        btn_frame = ttk.Frame(self)
        btn_frame.pack(fill="x", **pad)
        ttk.Button(btn_frame, text="미리보기 (Dry Run)", command=lambda: self._run(dry_run=True)).pack(side="left")
        ttk.Button(btn_frame, text="노트 삽입 실행", command=lambda: self._run(dry_run=False)).pack(side="left", padx=8)
        ttk.Button(btn_frame, text="방금 추출한 노트 사용", command=self._use_last_extracted).pack(side="right")

    def _browse_notes(self):
        path = filedialog.askopenfilename(
            title="슬라이드 노트 파일 선택", filetypes=[("Markdown / Text", "*.md *.txt"), ("모든 파일", "*.*")]
        )
        if path:
            self.notes_path.set(path)

    def _browse_pptx(self):
        path = filedialog.askopenfilename(
            title="원본 PPTX 파일 선택", filetypes=[("PowerPoint 파일", "*.pptx"), ("모든 파일", "*.*")]
        )
        if path:
            self.input_pptx_path.set(path)
            if not self.output_pptx_path.get():
                base, ext = os.path.splitext(path)
                self.output_pptx_path.set(f"{base}_노트삽입{ext}")

    def _browse_output(self):
        default = self.output_pptx_path.get() or "output_노트삽입.pptx"
        path = filedialog.asksaveasfilename(
            title="저장할 파일 선택",
            defaultextension=".pptx",
            initialfile=os.path.basename(default),
            initialdir=os.path.dirname(default) or None,
            filetypes=[("PowerPoint 파일", "*.pptx")],
        )
        if path:
            self.output_pptx_path.set(path)

    def _use_last_extracted(self):
        last = self.shared_state.get("last_extracted_file")
        if not last:
            messagebox.showinfo("안내", "먼저 '노트 추출' 탭에서 텍스트 파일로 저장하면 여기서 바로 불러올 수 있습니다.")
            return
        self.notes_path.set(last)
        self.log(f"'노트 추출' 탭에서 저장한 파일을 불러왔습니다: {last}")

    def _run(self, dry_run):
        notes_file = self.notes_path.get().strip()
        input_pptx = self.input_pptx_path.get().strip()
        output_pptx = self.output_pptx_path.get().strip()

        if not notes_file or not input_pptx:
            messagebox.showwarning("입력 확인", "노트 파일과 원본 PPTX 파일을 모두 선택하세요.")
            return
        if not dry_run and not output_pptx:
            messagebox.showwarning("입력 확인", "저장할 파일 경로를 지정하세요.")
            return
        if not dry_run and os.path.abspath(output_pptx) == os.path.abspath(input_pptx):
            messagebox.showwarning("입력 확인", "저장 경로가 원본 PPTX와 동일합니다. 다른 파일명으로 저장하세요.")
            return

        self.log("\n===== [노트 삽입] " + ("미리보기" if dry_run else "실행") + " 시작 =====")

        def worker():
            try:
                insert_slide_notes_from_file(notes_file, input_pptx, output_pptx, dry_run=dry_run, log=self.log)
            except Exception:
                self.log("예상치 못한 오류:\n" + traceback.format_exc())

        threading.Thread(target=worker, daemon=True).start()


class ExtractTab(ttk.Frame):
    def __init__(self, parent, log_fn, shared_state):
        super().__init__(parent)
        self.log = log_fn
        self.shared_state = shared_state

        self.input_pptx_path = tk.StringVar()
        self.output_path = tk.StringVar()
        self.fmt_var = tk.StringVar(value="markdown")  # "markdown" | "dash"
        self._auto_suggestion = ""

        self._build()

    def _build(self):
        pad = {"padx": 10, "pady": 6}
        form = ttk.Frame(self)
        form.pack(fill="x", **pad)
        form.columnconfigure(0, weight=1)

        ttk.Label(form, text="① PPTX 파일 (노트를 꺼낼 원본)").grid(row=0, column=0, columnspan=2, sticky="w")
        ttk.Entry(form, textvariable=self.input_pptx_path).grid(row=1, column=0, sticky="ew", padx=(0, 8))
        ttk.Button(form, text="찾아보기", command=self._browse_pptx).grid(row=1, column=1)

        fmt_frame = ttk.LabelFrame(self, text="② 슬라이드 구분 방식")
        fmt_frame.pack(fill="x", **pad)
        ttk.Radiobutton(
            fmt_frame,
            text="마크다운 형식  (## 슬라이드 01 / **슬라이드 노트 내용** — PPT 재삽입 시 가장 안전)",
            variable=self.fmt_var, value="markdown", command=self._on_fmt_change,
        ).pack(anchor="w", padx=8, pady=(6, 2))
        ttk.Radiobutton(
            fmt_frame,
            text="구분선 형식  (----- 로만 슬라이드 구분 — 읽고 편집하기 간단)",
            variable=self.fmt_var, value="dash", command=self._on_fmt_change,
        ).pack(anchor="w", padx=8, pady=(0, 6))

        form2 = ttk.Frame(self)
        form2.pack(fill="x", **pad)
        form2.columnconfigure(0, weight=1)
        ttk.Label(form2, text="③ 저장할 텍스트 파일").grid(row=0, column=0, columnspan=2, sticky="w")
        ttk.Entry(form2, textvariable=self.output_path).grid(row=1, column=0, sticky="ew", padx=(0, 8))
        ttk.Button(form2, text="다른 이름으로 저장", command=self._browse_output).grid(row=1, column=1)

        btn_frame = ttk.Frame(self)
        btn_frame.pack(fill="x", **pad)
        ttk.Button(btn_frame, text="텍스트로 추출", command=self._run_extract).pack(side="left")

    def _browse_pptx(self):
        path = filedialog.askopenfilename(
            title="PPTX 파일 선택", filetypes=[("PowerPoint 파일", "*.pptx"), ("모든 파일", "*.*")]
        )
        if path:
            self.input_pptx_path.set(path)
            self._suggest_output_path()

    def _on_fmt_change(self):
        self._suggest_output_path()

    def _suggest_output_path(self):
        pptx_path = self.input_pptx_path.get().strip()
        if not pptx_path:
            return
        current = self.output_path.get().strip()
        if current and current != self._auto_suggestion:
            return  # 사용자가 직접 수정한 값이면 덮어쓰지 않음
        p = Path(pptx_path)
        ext = ".md" if self.fmt_var.get() == "markdown" else ".txt"
        suggestion = str(p.with_name(f"{p.stem}_발표자노트{ext}"))
        self._auto_suggestion = suggestion
        self.output_path.set(suggestion)

    def _browse_output(self):
        default = self.output_path.get() or "발표자노트.txt"
        ext = ".md" if self.fmt_var.get() == "markdown" else ".txt"
        path = filedialog.asksaveasfilename(
            title="저장할 파일 선택",
            defaultextension=ext,
            initialfile=os.path.basename(default),
            initialdir=os.path.dirname(default) or None,
            filetypes=[("Markdown", "*.md"), ("Text", "*.txt"), ("모든 파일", "*.*")],
        )
        if path:
            self.output_path.set(path)

    def _run_extract(self):
        input_pptx = self.input_pptx_path.get().strip()
        output_path = self.output_path.get().strip()

        if not input_pptx:
            messagebox.showwarning("입력 확인", "PPTX 파일을 선택하세요.")
            return
        if not output_path:
            messagebox.showwarning("입력 확인", "저장할 파일 경로를 지정하세요.")
            return

        self.log("\n===== [노트 추출] 시작 =====")

        def worker():
            try:
                if not check_pptx_path(input_pptx, self.log):
                    return
                notes_data = extract_notes_from_pptx(input_pptx, self.log)

                if self.fmt_var.get() == "markdown":
                    text = format_notes_as_markdown(notes_data)
                else:
                    text = format_notes_as_dash(notes_data)

                with open(output_path, "w", encoding="utf-8") as f:
                    f.write(text)

                self.log(f"'{output_path}' 파일로 저장했습니다. (형식: {self.fmt_var.get()})")

                # 다른 탭에서 바로 이어서 쓸 수 있도록 공유 상태에 기록
                self.shared_state["last_extracted_file"] = output_path
                self.shared_state["last_extracted_notes"] = notes_data
                self.log("→ '노트 삽입' / '음성 변환' 탭에서 방금 추출한 결과를 바로 불러올 수 있습니다.")
            except Exception:
                self.log("예상치 못한 오류:\n" + traceback.format_exc())

        threading.Thread(target=worker, daemon=True).start()


class TTSTab(ttk.Frame):
    def __init__(self, parent, log_fn, shared_state):
        super().__init__(parent)
        self.log = log_fn
        self.shared_state = shared_state

        self.notes_source_path = tk.StringVar()
        self.api_key = tk.StringVar(value=os.environ.get("AZURE_SPEECH_KEY", ""))
        self.region_var = tk.StringVar(value=os.environ.get("AZURE_SPEECH_REGION", "koreacentral"))
        self.locale_var = tk.StringVar(value=AZURE_DEFAULT_LOCALE)
        self.voice_var = tk.StringVar(value="")
        self.range_var = tk.StringVar()
        self.output_dir = tk.StringVar()
        self.combine_var = tk.BooleanVar(value=False)
        self.combined_path = tk.StringVar()

        self._loaded_notes_data = None  # '방금 추출한 노트 사용'으로 불러온 경우 여기 채워짐

        self._build()

    def _build(self):
        pad = {"padx": 10, "pady": 6}

        src_frame = ttk.LabelFrame(self, text="① 노트 소스 (.md / .txt)")
        src_frame.pack(fill="x", **pad)
        src_frame.columnconfigure(0, weight=1)
        ttk.Entry(src_frame, textvariable=self.notes_source_path).grid(row=0, column=0, sticky="ew", padx=(8, 8), pady=8)
        ttk.Button(src_frame, text="파일에서 불러오기", command=self._browse_notes).grid(row=0, column=1, padx=(0, 4))
        ttk.Button(src_frame, text="방금 추출한 노트 사용", command=self._use_last_extracted).grid(row=0, column=2, padx=(0, 8))

        key_frame = ttk.LabelFrame(self, text="② Azure AI Speech 설정 (무료 F0 티어: 매달 뉴럴 TTS 50만 자)")
        key_frame.pack(fill="x", **pad)
        key_frame.columnconfigure(1, weight=1)
        key_frame.columnconfigure(3, weight=1)

        ttk.Label(key_frame, text="API 키").grid(row=0, column=0, sticky="w", padx=8, pady=(8, 4))
        ttk.Entry(key_frame, textvariable=self.api_key, show="*").grid(
            row=0, column=1, columnspan=3, sticky="ew", padx=(0, 8), pady=(8, 4)
        )

        ttk.Label(key_frame, text="리전").grid(row=1, column=0, sticky="w", padx=8, pady=4)
        ttk.Combobox(key_frame, textvariable=self.region_var, values=AZURE_COMMON_REGIONS, width=18).grid(
            row=1, column=1, sticky="w", pady=4
        )
        ttk.Label(key_frame, text="언어(로케일) 필터").grid(row=1, column=2, sticky="e", padx=(8, 4), pady=4)
        ttk.Entry(key_frame, textvariable=self.locale_var, width=14).grid(
            row=1, column=3, sticky="w", padx=(0, 8), pady=4
        )

        ttk.Label(
            key_frame,
            text="※ 리전은 Azure Portal에서 만든 Speech 리소스의 지역입니다(예: koreacentral). 로케일을 비우면 전체 언어를 불러옵니다.",
            foreground="#666666",
        ).grid(row=2, column=0, columnspan=4, sticky="w", padx=8)

        ttk.Label(key_frame, text="음성").grid(row=3, column=0, sticky="w", padx=8, pady=(8, 4))
        self.voice_combo = ttk.Combobox(key_frame, textvariable=self.voice_var, values=[], width=42)
        self.voice_combo.grid(row=3, column=1, columnspan=2, sticky="ew", padx=(0, 8), pady=(8, 4))
        ttk.Button(key_frame, text="음성 목록 불러오기", command=self._load_voices).grid(
            row=3, column=3, sticky="w", padx=(0, 8), pady=(8, 4)
        )
        ttk.Label(
            key_frame,
            text="※ 목록에서 고르거나, 음성 이름(예: ko-KR-SunHiNeural)을 알고 있으면 직접 입력해도 됩니다.",
            foreground="#666666",
        ).grid(row=4, column=0, columnspan=4, sticky="w", padx=8)

        ttk.Label(key_frame, text="슬라이드 범위 (예: 1-5,10 / 비워두면 전체)").grid(
            row=5, column=0, columnspan=2, sticky="w", padx=8, pady=(8, 8)
        )
        ttk.Entry(key_frame, textvariable=self.range_var).grid(
            row=5, column=2, columnspan=2, sticky="ew", padx=(0, 8), pady=(8, 8)
        )

        out_frame = ttk.LabelFrame(self, text="③ 출력")
        out_frame.pack(fill="x", **pad)
        out_frame.columnconfigure(0, weight=1)

        ttk.Label(out_frame, text="슬라이드별 mp3 저장 폴더").grid(row=0, column=0, sticky="w", padx=8, pady=(8, 0))
        ttk.Entry(out_frame, textvariable=self.output_dir).grid(row=1, column=0, sticky="ew", padx=(8, 8), pady=(0, 8))
        ttk.Button(out_frame, text="폴더 선택", command=self._browse_output_dir).grid(row=1, column=1, padx=(0, 8), pady=(0, 8))

        ttk.Checkbutton(
            out_frame, text="모든 슬라이드를 하나의 mp3로 합치기 (단순 결합)",
            variable=self.combine_var, command=self._toggle_combine,
        ).grid(row=2, column=0, sticky="w", padx=8)
        self.combined_entry = ttk.Entry(out_frame, textvariable=self.combined_path, state="disabled")
        self.combined_entry.grid(row=3, column=0, sticky="ew", padx=(8, 8), pady=(0, 8))
        self.combined_btn = ttk.Button(out_frame, text="합친 파일 경로 선택", command=self._browse_combined_path, state="disabled")
        self.combined_btn.grid(row=3, column=1, padx=(0, 8), pady=(0, 8))

        btn_frame = ttk.Frame(self)
        btn_frame.pack(fill="x", **pad)
        ttk.Button(btn_frame, text="음성 생성 시작", command=self._run_tts).pack(side="left")

    def _browse_notes(self):
        path = filedialog.askopenfilename(
            title="노트 파일 선택", filetypes=[("Markdown / Text", "*.md *.txt"), ("모든 파일", "*.*")]
        )
        if path:
            self.notes_source_path.set(path)
            self._loaded_notes_data = None
            if not self.output_dir.get():
                p = Path(path)
                self.output_dir.set(str(p.parent / f"{p.stem}_음성"))
                self.combined_path.set(str(p.parent / f"{p.stem}_전체.mp3"))

    def _use_last_extracted(self):
        last_file = self.shared_state.get("last_extracted_file")
        last_data = self.shared_state.get("last_extracted_notes")
        if not last_file:
            messagebox.showinfo("안내", "먼저 '노트 추출' 탭에서 텍스트 파일로 저장하면 여기서 바로 불러올 수 있습니다.")
            return
        self.notes_source_path.set(last_file)
        self._loaded_notes_data = last_data
        p = Path(last_file)
        if not self.output_dir.get():
            self.output_dir.set(str(p.parent / f"{p.stem}_음성"))
            self.combined_path.set(str(p.parent / f"{p.stem}_전체.mp3"))
        self.log(f"'노트 추출' 탭 결과를 불러왔습니다: {last_file}")

    def _load_voices(self):
        api_key = self.api_key.get().strip()
        region = self.region_var.get().strip()
        locale_filter = self.locale_var.get().strip()
        if not api_key or not region:
            messagebox.showwarning("입력 확인", "API 키와 리전을 먼저 입력하세요.")
            return

        self.log(f"음성 목록을 불러오는 중... (리전: {region}, 로케일: {locale_filter or '전체'})")

        def worker():
            voices = list_azure_voices(api_key, region, locale_filter, self.log)
            if not voices:
                return
            display = [f"{name} ({short_name})" for name, short_name in voices]

            def _apply():
                self.voice_combo.configure(values=display)
                if not self.voice_var.get():
                    self.voice_var.set(display[0])

            self.after(0, _apply)

        threading.Thread(target=worker, daemon=True).start()

    def _toggle_combine(self):
        state = "normal" if self.combine_var.get() else "disabled"
        self.combined_entry.configure(state=state)
        self.combined_btn.configure(state=state)

    def _browse_output_dir(self):
        path = filedialog.askdirectory(title="mp3 저장 폴더 선택")
        if path:
            self.output_dir.set(path)

    def _browse_combined_path(self):
        default = self.combined_path.get() or "전체.mp3"
        path = filedialog.asksaveasfilename(
            title="합친 mp3 파일 경로",
            defaultextension=".mp3",
            initialfile=os.path.basename(default),
            initialdir=os.path.dirname(default) or None,
            filetypes=[("MP3 파일", "*.mp3")],
        )
        if path:
            self.combined_path.set(path)

    def _run_tts(self):
        notes_source = self.notes_source_path.get().strip()
        api_key = self.api_key.get().strip()
        region = self.region_var.get().strip()
        output_dir = self.output_dir.get().strip()
        combine = self.combine_var.get()
        combined_path = self.combined_path.get().strip()

        if not notes_source:
            messagebox.showwarning("입력 확인", "노트 파일을 선택하거나 '방금 추출한 노트 사용'을 눌러주세요.")
            return
        if not api_key or not region:
            messagebox.showwarning("입력 확인", "Azure API 키와 리전을 모두 입력하세요.")
            return
        if not self.voice_var.get().strip():
            messagebox.showwarning("입력 확인", "음성을 선택하거나 이름을 입력하세요. ('음성 목록 불러오기' 버튼 이용 가능)")
            return
        if not output_dir:
            messagebox.showwarning("입력 확인", "mp3 저장 폴더를 지정하세요.")
            return
        if combine and not combined_path:
            messagebox.showwarning("입력 확인", "합친 mp3 파일의 저장 경로를 지정하세요.")
            return

        self.log("\n===== [음성 변환] 시작 =====")

        def worker():
            try:
                notes_data = self._loaded_notes_data
                if notes_data is None:
                    notes, order, fmt = load_notes(notes_source)
                    notes_data = [(n, notes[n]) for n in order]
                    self.log(f"노트 파일 형식: {fmt} / 총 {len(notes_data)}개 슬라이드")

                available = [n for n, _ in notes_data]
                slide_filter = parse_slide_range(self.range_var.get(), available)
                if slide_filter is not None and not slide_filter:
                    self.log("지정한 슬라이드 범위에 해당하는 슬라이드가 없습니다.")
                    return

                generate_audio_for_notes(
                    notes_data=notes_data,
                    api_key=api_key,
                    region=region,
                    voice_selection=self.voice_var.get(),
                    output_dir=output_dir,
                    combine=combine,
                    combined_path=combined_path,
                    slide_filter=slide_filter,
                    log=self.log,
                )
            except Exception:
                self.log("예상치 못한 오류:\n" + traceback.format_exc())

        threading.Thread(target=worker, daemon=True).start()


class PdfToPptTab(ttk.Frame):
    def __init__(self, parent, log_fn, shared_state):
        super().__init__(parent)
        self.log = log_fn
        self.shared_state = shared_state

        self.pdf_path = tk.StringVar()
        self.output_path = tk.StringVar()
        self.dpi = tk.IntVar(value=300)
        self.quality = tk.IntVar(value=95)
        self.status_var = tk.StringVar(value="PDF 파일을 선택하세요.")

        self._build()

    def _build(self):
        pad = {"padx": 10, "pady": 6}

        form = ttk.Frame(self)
        form.pack(fill="x", **pad)
        form.columnconfigure(0, weight=1)

        ttk.Label(form, text="① 변환할 PDF 파일").grid(row=0, column=0, columnspan=2, sticky="w")
        ttk.Entry(form, textvariable=self.pdf_path).grid(row=1, column=0, sticky="ew", padx=(0, 8))
        ttk.Button(form, text="찾아보기", command=self._browse_pdf).grid(row=1, column=1)

        ttk.Label(form, text="② 저장할 PPTX 위치").grid(row=2, column=0, columnspan=2, sticky="w", pady=(14, 0))
        ttk.Entry(form, textvariable=self.output_path).grid(row=3, column=0, sticky="ew", padx=(0, 8))
        ttk.Button(form, text="다른 이름으로 저장", command=self._browse_output).grid(row=3, column=1)

        opt_frame = ttk.LabelFrame(self, text="③ 화질 옵션")
        opt_frame.pack(fill="x", **pad)
        ttk.Label(opt_frame, text="해상도(DPI):").grid(row=0, column=0, padx=8, pady=8, sticky="w")
        ttk.Spinbox(opt_frame, from_=100, to=600, increment=50, textvariable=self.dpi, width=6).grid(
            row=0, column=1, padx=(0, 20), sticky="w"
        )
        ttk.Label(opt_frame, text="JPG 품질(1-100):").grid(row=0, column=2, padx=8, pady=8, sticky="w")
        ttk.Spinbox(opt_frame, from_=50, to=100, increment=5, textvariable=self.quality, width=6).grid(
            row=0, column=3, sticky="w"
        )
        ttk.Label(
            opt_frame,
            text="※ DPI가 높을수록 화질은 좋아지지만 변환 시간과 파일 용량이 늘어납니다. (권장: 300)",
            foreground="#666666",
        ).grid(row=1, column=0, columnspan=4, sticky="w", padx=8, pady=(0, 8))

        if not HAVE_PDF_LIBS:
            ttk.Label(
                self,
                text="⚠ pymupdf / pillow 가 설치되어 있지 않습니다. 터미널에서 'pip install pymupdf pillow' 실행 후 다시 시작하세요.",
                foreground="#b00020",
            ).pack(fill="x", padx=10, pady=(0, 6))

        btn_frame = ttk.Frame(self)
        btn_frame.pack(fill="x", **pad)
        self.convert_btn = ttk.Button(btn_frame, text="변환 시작", command=self._run_convert)
        self.convert_btn.pack(side="left")

        self.progress = ttk.Progressbar(self, orient="horizontal", mode="determinate")
        self.progress.pack(fill="x", padx=10, pady=(0, 4))

        ttk.Label(self, textvariable=self.status_var, foreground="#444").pack(padx=12, anchor="w")

    def _browse_pdf(self):
        path = filedialog.askopenfilename(title="PDF 파일 선택", filetypes=[("PDF 파일", "*.pdf")])
        if not path:
            return
        self.pdf_path.set(path)
        base = os.path.splitext(os.path.basename(path))[0]
        default_out = os.path.join(os.path.dirname(path), f"{base}.pptx")
        self.output_path.set(default_out)
        self.status_var.set("변환 시작 버튼을 누르세요.")

    def _browse_output(self):
        if not self.pdf_path.get():
            messagebox.showwarning("알림", "먼저 PDF 파일을 선택하세요.")
            return
        base = os.path.splitext(os.path.basename(self.pdf_path.get()))[0]
        path = filedialog.asksaveasfilename(
            title="PPTX 저장 위치",
            defaultextension=".pptx",
            initialfile=f"{base}.pptx",
            filetypes=[("PowerPoint 파일", "*.pptx")],
        )
        if path:
            self.output_path.set(path)

    def _run_convert(self):
        pdf_path = self.pdf_path.get().strip()
        output_path = self.output_path.get().strip()

        if not HAVE_PDF_LIBS:
            messagebox.showwarning(
                "라이브러리 필요", "이 기능을 쓰려면 터미널에서 'pip install pymupdf pillow' 를 먼저 실행하세요."
            )
            return
        if not pdf_path:
            messagebox.showwarning("입력 확인", "PDF 파일을 먼저 선택하세요.")
            return
        if not output_path:
            messagebox.showwarning("입력 확인", "저장할 PPTX 위치를 선택하세요.")
            return

        self.convert_btn.configure(state="disabled")
        self.progress.configure(value=0)
        self.log("\n===== [PDF → PPT] 변환 시작 =====")

        def progress_cb(done, total, msg):
            self.after(0, self._update_progress, done, total, msg)

        def worker():
            try:
                ok = convert_pdf_to_pptx(
                    pdf_path, output_path, self.dpi.get(), self.quality.get(),
                    progress_cb=progress_cb, log=self.log,
                )
                self.after(0, self._on_done, ok, output_path)
            except Exception:
                self.log("예상치 못한 오류:\n" + traceback.format_exc())
                self.after(0, self._on_done, False, output_path)

        threading.Thread(target=worker, daemon=True).start()

    def _update_progress(self, done, total, msg):
        self.progress.configure(maximum=total, value=done)
        self.status_var.set(msg)
        self.log(msg)

    def _on_done(self, ok, output_path):
        self.convert_btn.configure(state="normal")
        if not ok:
            self.status_var.set("오류 발생 (아래 로그 확인)")
            return
        self.status_var.set(f"완료: {output_path}")
        if messagebox.askyesno("완료", "변환이 완료되었습니다.\n저장된 폴더를 여시겠습니까?"):
            open_in_file_manager(output_path)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("PPT 발표자 노트 & 변환 도구 모음")
        self.geometry("860x820")
        self.minsize(760, 640)

        self.shared_state = {"last_extracted_file": None, "last_extracted_notes": None}

        notebook = ttk.Notebook(self)
        notebook.pack(fill="both", expand=True, padx=8, pady=(8, 0))

        self.insert_tab = InsertTab(notebook, self._log, self.shared_state)
        self.extract_tab = ExtractTab(notebook, self._log, self.shared_state)
        self.tts_tab = TTSTab(notebook, self._log, self.shared_state)
        self.pdf_tab = PdfToPptTab(notebook, self._log, self.shared_state)

        notebook.add(self.insert_tab, text="① 노트 삽입 (텍스트→PPT)")
        notebook.add(self.extract_tab, text="② 노트 추출 (PPT→텍스트)")
        notebook.add(self.tts_tab, text="③ 음성 변환 (텍스트→MP3)")
        notebook.add(self.pdf_tab, text="④ PDF → PPT (이미지 변환)")

        log_frame = ttk.LabelFrame(self, text="진행 로그")
        log_frame.pack(fill="both", expand=True, padx=8, pady=8)

        toolbar = ttk.Frame(log_frame)
        toolbar.pack(fill="x")
        ttk.Button(toolbar, text="지우기", command=self._clear_log).pack(side="right", padx=4, pady=4)

        self.log_text = scrolledtext.ScrolledText(log_frame, wrap="word", state="disabled", font=("Consolas", 10))
        self.log_text.pack(fill="both", expand=True, padx=4, pady=(0, 4))

    def _log(self, message):
        def _append():
            self.log_text.configure(state="normal")
            self.log_text.insert("end", str(message) + "\n")
            self.log_text.see("end")
            self.log_text.configure(state="disabled")
        self.after(0, _append)

    def _clear_log(self):
        self.log_text.configure(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.configure(state="disabled")


if __name__ == "__main__":
    app = App()
    app.mainloop()