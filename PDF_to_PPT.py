"""
PDF -> 고화질 JPG -> PPT 변환 GUI

PDF를 선택하면 각 페이지를 고해상도 JPG로 렌더링한 뒤,
슬라이드 전체에 꽉 채운 PPTX로 저장합니다. (이미지 삽입 방식이라
폰트 깨짐/레이아웃 깨짐이 없습니다)

필요 라이브러리 설치:
    pip install pymupdf python-pptx pillow

실행:
    python pdf_to_ppt_gui.py
"""

import os
import platform
import subprocess
import sys
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    import fitz  # PyMuPDF
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Emu
except ImportError as e:
    print("필요한 라이브러리가 설치되어 있지 않습니다.")
    print("아래 명령어로 먼저 설치해주세요:\n")
    print("    pip install pymupdf python-pptx pillow\n")
    print(f"(오류 내용: {e})")
    sys.exit(1)

EMU_PER_PX = 9525  # 96dpi 기준 1px = 9525 EMU


def pdf_to_images(pdf_path, out_dir, dpi, quality, progress_cb=None):
    doc = fitz.open(pdf_path)
    zoom = dpi / 72
    matrix = fitz.Matrix(zoom, zoom)
    total = len(doc)
    image_paths = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=matrix, alpha=False)
        img_path = os.path.join(out_dir, f"page_{i + 1:03d}.jpg")
        pix.save(img_path, jpg_quality=quality)
        image_paths.append(img_path)
        if progress_cb:
            progress_cb(i + 1, total, f"[JPG] {i + 1}/{total} 페이지 저장 완료 ({pix.width}x{pix.height}px)")
    doc.close()
    return image_paths


def images_to_ppt(image_paths, output_pptx, progress_cb=None):
    if not image_paths:
        raise ValueError("변환할 이미지가 없습니다.")

    with Image.open(image_paths[0]) as im:
        px_w, px_h = im.size

    prs = Presentation()
    prs.slide_width = Emu(px_w * EMU_PER_PX)
    prs.slide_height = Emu(px_h * EMU_PER_PX)
    blank_layout = prs.slide_layouts[6]

    total = len(image_paths)
    for i, img_path in enumerate(image_paths):
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(img_path, left=0, top=0, width=prs.slide_width, height=prs.slide_height)
        if progress_cb:
            progress_cb(i + 1, total, f"[PPT] {i + 1}/{total} 슬라이드 생성 완료")

    prs.save(output_pptx)


def open_in_file_manager(path):
    folder = os.path.dirname(path) or "."
    system = platform.system()
    try:
        if system == "Windows":
            os.startfile(folder)
        elif system == "Darwin":
            subprocess.run(["open", folder])
        else:
            subprocess.run(["xdg-open", folder])
    except Exception:
        pass


class PdfToPptApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("PDF → PPT 변환기 (고화질 이미지)")
        self.geometry("580x460")
        self.resizable(False, False)

        self.pdf_path = tk.StringVar()
        self.output_path = tk.StringVar()
        self.dpi = tk.IntVar(value=300)
        self.quality = tk.IntVar(value=95)
        self.status = tk.StringVar(value="PDF 파일을 선택하세요.")

        self._build_ui()

    def _build_ui(self):
        pad = {"padx": 10, "pady": 6}

        frame1 = ttk.LabelFrame(self, text="1. 변환할 PDF")
        frame1.pack(fill="x", **pad)
        ttk.Entry(frame1, textvariable=self.pdf_path, state="readonly").pack(
            side="left", fill="x", expand=True, padx=(8, 4), pady=8
        )
        ttk.Button(frame1, text="찾아보기...", command=self.browse_pdf).pack(side="left", padx=(0, 8))

        frame2 = ttk.LabelFrame(self, text="2. 저장할 PPTX 위치")
        frame2.pack(fill="x", **pad)
        ttk.Entry(frame2, textvariable=self.output_path, state="readonly").pack(
            side="left", fill="x", expand=True, padx=(8, 4), pady=8
        )
        ttk.Button(frame2, text="다른 이름으로...", command=self.browse_output).pack(side="left", padx=(0, 8))

        frame3 = ttk.LabelFrame(self, text="3. 화질 옵션")
        frame3.pack(fill="x", **pad)
        ttk.Label(frame3, text="해상도(DPI):").grid(row=0, column=0, padx=8, pady=8, sticky="w")
        ttk.Spinbox(frame3, from_=100, to=600, increment=50, textvariable=self.dpi, width=6).grid(
            row=0, column=1, padx=(0, 20), sticky="w"
        )
        ttk.Label(frame3, text="JPG 품질(1-100):").grid(row=0, column=2, padx=8, pady=8, sticky="w")
        ttk.Spinbox(frame3, from_=50, to=100, increment=5, textvariable=self.quality, width=6).grid(
            row=0, column=3, sticky="w"
        )

        self.convert_btn = ttk.Button(self, text="변환 시작", command=self.start_convert)
        self.convert_btn.pack(pady=(6, 4))

        self.progress = ttk.Progressbar(self, orient="horizontal", mode="determinate", length=540)
        self.progress.pack(padx=10, pady=4)

        ttk.Label(self, textvariable=self.status, foreground="#444").pack(padx=12, anchor="w")

        frame4 = ttk.LabelFrame(self, text="진행 로그")
        frame4.pack(fill="both", expand=True, **pad)
        self.log = tk.Text(frame4, height=9, state="disabled", wrap="word")
        self.log.pack(fill="both", expand=True, padx=8, pady=8)

    def browse_pdf(self):
        path = filedialog.askopenfilename(title="PDF 파일 선택", filetypes=[("PDF 파일", "*.pdf")])
        if not path:
            return
        self.pdf_path.set(path)
        base = os.path.splitext(os.path.basename(path))[0]
        default_out = os.path.join(os.path.dirname(path), f"{base}.pptx")
        self.output_path.set(default_out)
        self.status.set("변환 시작 버튼을 누르세요.")

    def browse_output(self):
        if not self.pdf_path.get():
            messagebox.showwarning("알림", "먼저 PDF 파일을 선택하세요.")
            return
        base = os.path.splitext(os.path.basename(self.pdf_path.get()))[0]
        path = filedialog.asksaveasfilename(
            title="PPTX 저장 위치",
            defaultextension=".pptx",
            initialfile=f"{base}.pptx",
            filetypes=[("PowerPoint 파일", "*.pptx")],
        )
        if path:
            self.output_path.set(path)

    def log_write(self, text):
        self.log.configure(state="normal")
        self.log.insert("end", text + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def start_convert(self):
        pdf_path = self.pdf_path.get()
        output_path = self.output_path.get()
        if not pdf_path:
            messagebox.showwarning("알림", "PDF 파일을 먼저 선택하세요.")
            return
        if not output_path:
            messagebox.showwarning("알림", "저장할 PPTX 위치를 선택하세요.")
            return

        self.convert_btn.configure(state="disabled")
        self.progress["value"] = 0
        self.log.configure(state="normal")
        self.log.delete("1.0", "end")
        self.log.configure(state="disabled")

        threading.Thread(target=self._run_conversion, args=(pdf_path, output_path), daemon=True).start()

    def _run_conversion(self, pdf_path, output_path):
        try:
            base_name = os.path.splitext(os.path.basename(pdf_path))[0]
            out_dir = os.path.join(os.path.dirname(output_path) or ".", f"{base_name}_slides")
            os.makedirs(out_dir, exist_ok=True)

            def jpg_progress(done, total, msg):
                self.after(0, self._update_progress, done, total, msg)

            def ppt_progress(done, total, msg):
                self.after(0, self.log_write, msg)

            self.after(0, self.status.set, "PDF → JPG 변환 중...")
            image_paths = pdf_to_images(
                pdf_path, out_dir, self.dpi.get(), self.quality.get(), progress_cb=jpg_progress
            )

            self.after(0, self.status.set, "PPT 생성 중...")
            images_to_ppt(image_paths, output_path, progress_cb=ppt_progress)

            self.after(0, self._on_done, output_path)
        except Exception as e:
            self.after(0, self._on_error, str(e))

    def _update_progress(self, done, total, msg):
        self.progress["maximum"] = total
        self.progress["value"] = done
        self.log_write(msg)

    def _on_done(self, output_path):
        self.status.set(f"완료: {output_path}")
        self.log_write(f"\n✅ 변환 완료: {output_path}")
        self.convert_btn.configure(state="normal")
        if messagebox.askyesno("완료", "변환이 완료되었습니다.\n저장된 폴더를 여시겠습니까?"):
            open_in_file_manager(output_path)

    def _on_error(self, message):
        self.status.set("오류 발생")
        self.log_write(f"\n❌ 오류: {message}")
        self.convert_btn.configure(state="normal")
        messagebox.showerror("오류", message)


if __name__ == "__main__":
    app = PdfToPptApp()
    app.mainloop()