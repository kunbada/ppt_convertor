"""PPTX -> PDF / 슬라이드별 이미지 / 이미지형 PPTX 변환 (Windows + PowerPoint 필요)

사용법:
  python ppt_export.py 파일.pptx pdf            # 파일.pdf
  python ppt_export.py 파일.pptx images         # 파일_이미지/슬라이드_001.png ...
  python ppt_export.py 파일.pptx flat           # 파일_이미지형.pptx (각 슬라이드가 이미지 한 장, 노트 유지)
  python ppt_export.py 파일.pptx all            # 위 셋 모두
옵션: --width 1920 (이미지 가로 픽셀, 기본 1920)  --fmt png|jpg

필요: pip install pywin32 python-pptx
"""
import argparse
import os
import sys
import tempfile

PP_SAVE_AS_PDF = 32


def export(src, mode, width, fmt):
    import win32com.client
    from pptx import Presentation

    src = os.path.abspath(src)
    base = os.path.splitext(src)[0]
    ppt = win32com.client.Dispatch("PowerPoint.Application")
    pres = ppt.Presentations.Open(src, WithWindow=False, ReadOnly=True)
    try:
        if mode in ("pdf", "all"):
            pres.SaveAs(base + ".pdf", PP_SAVE_AS_PDF)
            print("✓", base + ".pdf")

        if mode in ("images", "flat", "all"):
            img_dir = base + "_이미지" if mode != "flat" else tempfile.mkdtemp()
            os.makedirs(img_dir, exist_ok=True)
            w = int(width)
            h = round(w * pres.PageSetup.SlideHeight / pres.PageSetup.SlideWidth)
            paths = []
            for i, slide in enumerate(pres.Slides, 1):
                p = os.path.join(img_dir, f"슬라이드_{i:03d}.{fmt}")
                slide.Export(p, "PNG" if fmt == "png" else "JPG", w, h)
                paths.append(p)
            if mode != "flat":
                print(f"✓ {len(paths)}장 → {img_dir}")

            if mode in ("flat", "all"):
                # 원본을 복제해 각 슬라이드 내용을 지우고 이미지 한 장으로 교체 (노트·크기 유지)
                out = base + "_이미지형.pptx"
                prs = Presentation(src)
                for slide, p in zip(prs.slides, paths):
                    for shp in list(slide.shapes):
                        shp._element.getparent().remove(shp._element)
                    slide.shapes.add_picture(p, 0, 0, prs.slide_width, prs.slide_height)
                prs.save(out)
                print("✓", out)
    finally:
        pres.Close()
        ppt.Quit()


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("mode", choices=["pdf", "images", "flat", "all"])
    ap.add_argument("--width", default=1920)
    ap.add_argument("--fmt", default="png", choices=["png", "jpg"])
    a = ap.parse_args()
    if sys.platform != "win32":
        sys.exit("Windows + PowerPoint 환경에서만 동작합니다. (macOS/Linux는 LibreOffice를 사용하세요)")
    export(a.pptx, a.mode, a.width, a.fmt)
